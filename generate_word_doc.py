import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_proposal_docx(filename="Plan_Sistema_Creditos_y_Control_Pagos_VOX.docx"):
    doc = Document()

    # Configuración de márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colores de marca VOX
    HEX_PRIMARY = "0D1117"     # Dark slate / Header
    HEX_ORANGE = "FF5500"      # VOX Accent Fox Orange
    HEX_DARK_BLUE = "1F2937"    # Subtitles / Cards
    HEX_MUTED = "4B5563"        # Body text
    HEX_LIGHT_BG = "F3F4F6"     # Table headers / Light boxes
    HEX_BORDER = "E5E7EB"

    COLOR_PRIMARY = RGBColor(13, 17, 23)
    COLOR_ORANGE = RGBColor(255, 85, 0)
    COLOR_DARK = RGBColor(31, 41, 55)
    COLOR_MUTED = RGBColor(75, 85, 99)

    # Helper para sombreado de celdas
    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    # Helper para bordes de tabla
    def set_table_borders(table, color="D1D5DB"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # --- ENCABEZADO / TÍTULO PRINCIPAL ---
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_brand = p_title.add_run("VOX BUSINESS DEVELOPER\n")
    run_brand.font.size = Pt(12)
    run_brand.font.bold = True
    run_brand.font.color.rgb = COLOR_ORANGE

    run_title = p_title.add_run("Propuesta Técnica: Sistema de Créditos, Cobranza y Bloqueo de Sitios Web (Killswitch)")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Plan Operativo de Financiamiento (8 Meses x $1,000 MXN), Control Administrativo y Protección Tecnológica del Servicio")
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_MUTED
    run_sub.font.italic = True

    # Divisor visual
    p_hr = doc.add_paragraph()
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(16)
    r_hr = p_hr.add_run("―" * 55)
    r_hr.font.color.rgb = RGBColor(229, 231, 235)

    # --- 1. RESUMEN EJECUTIVO ---
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Resumen Ejecutivo")
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY
    h1.paragraph_format.space_after = Pt(6)

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.space_after = Pt(10)
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.add_run(
        "Con el objetivo de facilitar el acceso a sitios web de alto nivel para empresas y emprendedores, VOX implementará un esquema de financiamiento propio consistente en "
    )
    r_bold1 = p_exec.add_run("8 mensualidades de $1,000.00 MXN (Total: $8,000.00 MXN)")
    r_bold1.bold = True
    p_exec.add_run(
        ". Para garantizar la rentabilidad, certidumbre en el cobro y minimizar el riesgo de morosidad, se diseñará un sistema tecnológico integral con "
    )
    r_bold2 = p_exec.add_run("cobranza manual eficiente, registro de pagos y un mecanismo de desactivación inmediata (Killswitch) ")
    r_bold2.bold = True
    p_exec.add_run(
        "que suspende la página web del cliente de forma automatizada y elegante en caso de impago."
    )

    # --- 2. PILARES DEL SISTEMA ---
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Los Tres Pilares de la Solución")
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY
    h2.paragraph_format.space_after = Pt(6)

    # Tarjeta 1
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_after = Pt(4)
    r_p1_title = p_p1.add_run("A. Script de Verificación y Bloqueo (Killswitch Embebible)\n")
    r_p1_title.bold = True
    r_p1_title.font.size = Pt(11)
    r_p1_title.font.color.rgb = COLOR_ORANGE
    p_p1.add_run(
        "• Snippet de código ultraligero (<3KB) insertado en el encabezado de las webs de los clientes.\n"
        "• Al cargar la web, consulta en milisegundos el estado de la licencia en la nube de VOX.\n"
        "• Si el cliente está 'Al Corriente': la web carga con velocidad óptima y 100% transparente.\n"
        "• Si el cliente está 'Suspendido': se despliega un telón de bloqueo institucional VOX indicando que el sitio se encuentra temporalmente inactivo por gestión administrativa, con botón directo a WhatsApp de VOX."
    )

    # Tarjeta 2
    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_after = Pt(4)
    r_p2_title = p_p2.add_run("B. Módulo de Créditos y Cobranza en Dashboard VOX (/admin)\n")
    r_p2_title.bold = True
    r_p2_title.font.size = Pt(11)
    r_p2_title.font.color.rgb = COLOR_ORANGE
    p_p2.add_run(
        "• Pestaña dedicada 'Créditos & Sitios' dentro del panel de administración actual.\n"
        "• Monitor en tiempo real: Clientes activos, cuotas cobradas en el mes y sitios en mora.\n"
        "• Barra de progreso visual por cliente (ej. '3 de 8 Meses Pagados - $3,000 / $8,000 MXN').\n"
        "• Interruptor Killswitch manual (1-Click) para suspender o reactivar un sitio web en segundos."
    )

    # Tarjeta 3
    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_after = Pt(12)
    r_p3_title = p_p3.add_run("C. Flujo Operativo de Cobro Manual Simplificado\n")
    r_p3_title.bold = True
    r_p3_title.font.size = Pt(11)
    r_p3_title.font.color.rgb = COLOR_ORANGE
    p_p3.add_run(
        "• Registro de pagos en 1 solo clic tras recibir transferencias SPEI, OXXO o efectivo.\n"
        "• Historial de transacciones con fecha, método de pago y folio de referencia.\n"
        "• Botón de 'Recordatorio por WhatsApp' que redacta un mensaje automático con el saldo y fecha de corte del cliente.\n"
        "• Liberación permanente (Liquidación): Al completar el mes 8/8, la web queda liberada de forma definitiva."
    )

    # --- 3. TABLA COMPARATIVA DE ESTADOS DEL CLIENTE ---
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Estados de Licencia y Reglas del Negocio")
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_PRIMARY
    h3.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    headers = ["Estado del Sitio", "Condición / Regla de Tiempo", "Efecto Visual en la Web"]
    for i, h_text in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, HEX_PRIMARY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    rows_data = [
        ("Activo (Al Corriente)", "Mensualidad pagada dentro de la fecha de corte.", "Sitio 100% funcional y visible para sus clientes."),
        ("En Gracia (Tolerancia)", "1 a 3 días posteriores al corte sin registrar pago.", "Sitio visible. Alerta preventiva en Dashboard para recordatorio por WhatsApp."),
        ("Suspendido (Mora)", "4+ días vencido o forzado manualmente por el admin.", "Pantalla bloqueada con aviso elegante VOX y botón de pago/contacto."),
        ("Liquidado (8/8 Pagos)", "Completó los 8 meses acordados ($8,000 MXN).", "Sitio liberado permanentemente de cualquier restricción.")
    ]

    for row_idx, (st, cond, eff) in enumerate(rows_data, start=1):
        c0 = table.cell(row_idx, 0)
        c1 = table.cell(row_idx, 1)
        c2 = table.cell(row_idx, 2)
        
        # Fondo alternado
        bg = HEX_LIGHT_BG if row_idx % 2 == 0 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_background(c2, bg)

        p0 = c0.paragraphs[0]; p0.paragraph_format.space_before = Pt(4); p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(st); r0.bold = True; r0.font.size = Pt(9.5)
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_before = Pt(4); p1.paragraph_format.space_after = Pt(4)
        r1 = p1.add_run(cond); r1.font.size = Pt(9.5)

        p2 = c2.paragraphs[0]; p2.paragraph_format.space_before = Pt(4); p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(eff); r2.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- 4. CLÁUSULA LEGAL SUGERIDA PARA CONTRATOS ---
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Respaldo Legal (Cláusula de Reserva de Dominio)")
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY
    h4.paragraph_format.space_after = Pt(6)

    p_legal = doc.add_paragraph()
    p_legal.paragraph_format.line_spacing = 1.15
    p_legal.paragraph_format.space_after = Pt(12)
    p_legal.add_run(
        "Para total transparencia con el cliente, en la cotización o contrato de servicio se incluirá la siguiente especificación:\n\n"
    )
    r_quote = p_legal.add_run(
        '«El desarrollo web se entrega bajo esquema de financiamiento directo en 8 mensualidades sucesivas de $1,000.00 MXN con reserva de dominio del código y hosting. El servicio de visualización y disponibilidad del sitio web está condicionado al cumplimiento puntual de cada cuota. En caso de presentar un retraso superior a los días de gracia establecidos, el sitio web entrará en suspensión técnica temporal hasta la liquidación de la mensualidad vencida, sin responsabilidad para VOX Business Developer por pérdidas comerciales derivadas de dicha suspensión.»'
    )
    r_quote.italic = True
    r_quote.font.size = Pt(10)
    r_quote.font.color.rgb = COLOR_DARK

    # --- 5. PLAN DE IMPLEMENTACIÓN Y ENTREGABLES ---
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Cronograma de Entrega y Fases Técnicas")
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = COLOR_PRIMARY
    h5.paragraph_format.space_after = Pt(6)

    p_fases = doc.add_paragraph()
    p_fases.paragraph_format.space_after = Pt(4)
    p_fases.add_run(
        "1. Base de Datos en Supabase: Creación de tablas de créditos, historial de pagos y endpoint de validación seguro.\n"
        "2. Script Killswitch (vox-license.js): Código embebible con overlay corporativo de suspensión y botón de contacto.\n"
        "3. Interfaz en Dashboard (/admin): Pestaña de Créditos, registro de pagos SPEI/Efectivo, barras de avance y switch de suspensión.\n"
        "4. Generador de Snippets: Botón para copiar el script listo con el token del cliente para pegar en su página web.\n"
        "5. Pruebas de Despliegue: Simulación de corte, suspensión en tiempo real y reactivación instantánea."
    )

    # --- FIRMA / PIE DE PÁGINA ---
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(24)
    r_sign1 = p_sign.add_run("Preparado para Dirección General\n")
    r_sign1.font.size = Pt(10)
    r_sign1.font.color.rgb = COLOR_MUTED
    r_sign2 = p_sign.add_run("VOX Business Developer — Área de Desarrollo Tecnológico & Operaciones")
    r_sign2.font.size = Pt(10)
    r_sign2.bold = True
    r_sign2.font.color.rgb = COLOR_PRIMARY

    # Guardar documento
    doc.save(filename)
    print(f"Documento creado exitosamente: {filename}")

if __name__ == "__main__":
    create_proposal_docx()
