import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_vox_architecture_manual():
    doc = docx.Document()

    # Configuración de página y márgenes
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Paleta de Colores Corporativos VOX
    COLOR_PRIMARY = RGBColor(255, 106, 0)      # VOX Orange #FF6A00
    COLOR_DARK = RGBColor(22, 26, 31)          # Dark Slate #161A1F
    COLOR_MUTED = RGBColor(100, 110, 125)      # Slate Muted #646E7D
    COLOR_SUCCESS = RGBColor(39, 174, 96)      # Emerald Green #27AE60
    COLOR_ALERT = RGBColor(211, 84, 0)         # Amber Warning #D35400
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_CODE_BG = 'F4F6F8'

    def set_cell_background(cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_section_header(title_text, num_badge=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        
        if num_badge:
            r_num = p.add_run(f"{num_badge} ")
            r_num.font.name = 'Segoe UI'
            r_num.font.size = Pt(13)
            r_num.font.bold = True
            r_num.font.color.rgb = COLOR_PRIMARY

        r_title = p.add_run(title_text)
        r_title.font.name = 'Segoe UI'
        r_title.font.size = Pt(13)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_DARK

    def add_callout(text, prefix="PRINCIPIO ARQUITECTÓNICO:", border_color="FF6A00", bg_color="FFF8F3"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, bg_color)
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        
        # Left border highlight
        tcPr = c._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if prefix:
            r_p = p.add_run(f"{prefix}\n")
            r_p.font.name = 'Segoe UI'
            r_p.font.size = Pt(9.5)
            r_p.font.bold = True
            r_p.font.color.rgb = COLOR_PRIMARY
        
        r_t = p.add_run(text)
        r_t.font.name = 'Segoe UI'
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_DARK
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # -------------------------------------------------------------
    # 1. ENCABEZADO EJECUTIVO / BANNER DE REPORTE TÉCNICO
    # -------------------------------------------------------------
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cell = header_table.cell(0, 0)
    set_cell_background(header_cell, '161A1F')
    set_cell_margins(header_cell, top=200, bottom=200, left=240, right=240)

    hp = header_cell.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_sub = hp.add_run('VOX TECH & WEB ARCHITECTURE — SISTEMA DE INGENIERÍA\n')
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_PRIMARY

    r_title = hp.add_run('MANUAL DE ARQUITECTURA DE DESARROLLO WEB Y ESTÁNDAR DE SISTEMAS\n')
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(15.5)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_WHITE

    r_desc = hp.add_run('Guía Maestra de Capas, Dev/Ops, Calidad de Código (QA), Seguridad Ampliada, Monitoreo y Reglas de Homologación')
    r_desc.font.name = 'Segoe UI'
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = RGBColor(200, 205, 215)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Ficha Técnica
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    meta_data = [
        [('Ecosistema', 'VOX Business Developer Multi-Service Web'),
         ('Estrategia Git', 'GitFlow (main, staging, feature/*)'),
         ('Versión del Estándar', 'v2.0 — Producción 2026')],
        [('Despliegue & CI/CD', 'Vercel Serverless / Automated QA Gate'),
         ('Seguridad & Datos', 'Supabase PostgreSQL + RLS + API Shield'),
         ('Observabilidad', 'Sentry / Log Centralizado & Telemetría')]
    ]

    for r_idx, row in enumerate(meta_data):
        for c_idx, (label, val) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_background(cell, 'F8F9FA')
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r_l = p.add_run(f'{label}:\n')
            r_l.font.name = 'Segoe UI'
            r_l.font.size = Pt(8)
            r_l.font.bold = True
            r_l.font.color.rgb = COLOR_MUTED
            r_v = p.add_run(val)
            r_v.font.name = 'Segoe UI'
            r_v.font.size = Pt(9)
            r_v.font.bold = True
            r_v.font.color.rgb = COLOR_DARK

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

    # Principio Arquitectónico
    add_callout(
        "Toda página o aplicación desarrollada por VOX debe construirse bajo una arquitectura modular, escalable y desacoplada. "
        "Ninguna página web es una isla de código independiente: todas comparten un Núcleo Central de Utilidades (Core Engine), "
        "módulos funcionales autónomos y reutilizables, una capa blindada de seguridad y captura de datos (lead-service.js con Cloudflare Turnstile / reCAPTCHA), "
        "y un pipeline automatizado de CI/CD que garantiza cero degradación técnica a lo largo del tiempo.",
        prefix="PRINCIPIO ARQUITECTÓNICO DE DESARROLLO:"
    )

    # -------------------------------------------------------------
    # 1. VISIÓN DEL SISTEMA WEB INTEGRADO VOX
    # -------------------------------------------------------------
    add_section_header("Visión del Sistema Web Integrado VOX", "1.")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "El estándar de desarrollo web de VOX establece que cada sitio o landing page de servicio "
        "(desarrollo de software, consultoría jurídica, asesoría contable, arquitectura comercial, marketing, etc.) "
        "debe funcionar como una extensión natural de un único ecosistema tecnológico.\n\n"
        "Se elimina de forma tajante el código 'espagueti', las librerías duplicadas y los formularios desconectados. "
        "En su lugar, se implementa una arquitectura por capas estrictamente delimitadas que asegura velocidad de carga ultrarrápida, "
        "consistencia visual bajo un Design System unificado, seguridad de datos blindada contra spam/ataques, observabilidad en tiempo real "
        "y un mantenimiento ágil a largo plazo."
    )
    r.font.name = 'Segoe UI'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # 2. LAS CINCO CAPAS DE LA ARQUITECTURA WEB VOX
    # -------------------------------------------------------------
    add_section_header("Las Cinco Capas de la Arquitectura Web VOX", "2.")

    layers = [
        ("Capa 1: Núcleo y Fundamentos Compartidos (Core Layer)",
         "Contiene los estilos globales (tokens CSS, variables de diseño, tipografías y reset), librerías base, constantes globales, matriz unificada de variables de entorno (.env) y configuración centralizada de conexión (Supabase client, utilidades de red seguras)."),
        ("Capa 2: Módulos de Componentes Reutilizables (Component Layer)",
         "Bloques de UI estandarizados, accesibles y autónomos: Header/Navbar inteligente con estado activo, Hero sections con call-to-actions dinámicos, Cards de servicios con micro-interacciones, Formularios interactivos con validación instantánea y protección anti-bots (Turnstile/reCAPTCHA v3), Modales modulares, Acordeones FAQ y Footers corporativos."),
        ("Capa 3: Lógica de Negocio, Servicios y Observabilidad (Service & Business Logic Layer)",
         "Módulos JavaScript independientes y desacoplados del DOM: captura, validación y sanitización universal de leads (lead-service.js), gestión de licencias y créditos (vox-license.js), telemetría, analítica de eventos y monitoreo centralizado de errores (Sentry / LogSnag / Error Boundary)."),
        ("Capa 4: Seguridad, Autenticación y Persistencia Backend (Security & Data Layer)",
         "Integración con backend serverless seguro (/api/*.js) protegido con CORS restrictivo, Rate Limiting por IP (Upstash Redis / Token Bucket), mitigación anti-DDoS, políticas RLS (Row Level Security) en PostgreSQL/Supabase, control de roles (RBAC) y sanitización estricta."),
        ("Capa 5: Experiencia de Usuario, Cumplimiento y Rendimiento (UX & Performance Layer)",
         "Micro-interacciones optimizadas a 60 FPS, lazy loading nativo de recursos multimedia, versionado de assets para cache-busting inteligente, accesibilidad semántica WCAG/HTML5 (SEO técnico con único <h1> estructurado), Core Web Vitals (< 1.5s LCP), y gestión de privacidad/cookies (Banner de consentimiento GDPR/LGPD compliant).")
    ]

    for title, desc in layers:
        p_l = doc.add_paragraph()
        p_l.paragraph_format.space_before = Pt(3)
        p_l.paragraph_format.space_after = Pt(4)
        
        r_bullet = p_l.add_run("▪ ")
        r_bullet.font.color.rgb = COLOR_PRIMARY
        r_bullet.font.bold = True
        
        r_lt = p_l.add_run(f"{title}: ")
        r_lt.font.name = 'Segoe UI'
        r_lt.font.size = Pt(9.5)
        r_lt.font.bold = True
        r_lt.font.color.rgb = COLOR_DARK

        r_ld = p_l.add_run(desc)
        r_ld.font.name = 'Segoe UI'
        r_ld.font.size = Pt(9.5)
        r_ld.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # 3. ESTRUCTURA ESTÁNDAR DE DIRECTORIOS Y ARCHIVOS
    # -------------------------------------------------------------
    add_section_header("Estructura Estándar de Archivos del Proyecto", "3.")
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run("Toda implementación web bajo el estándar VOX debe respetar la siguiente organización jerárquica de archivos:")
    r_t.font.name = 'Segoe UI'
    r_t.font.size = Pt(9.5)

    files_table = doc.add_table(rows=1, cols=3)
    files_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = files_table.rows[0].cells
    hdr_titles = ['Directorio / Archivo', 'Capa / Naturaleza', 'Propósito y Regla de Uso']
    for i, title in enumerate(hdr_titles):
        set_cell_background(hdr_cells[i], '161A1F')
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_WHITE

    files_data = [
        ('/assets/css/ (styles.css, vars.css)', 'Capa 1: Core', 'Tokens globales de color, tipografía, espaciado y utilidades base. Prohibido estilos inline o CSS aislado no tokenizado.'),
        ('.env.example / .env.local', 'Capa 1: Core / Envs', 'Plantilla y variables locales seguras. Nunca se commitean secretos ni API Keys privadas al repositorio.'),
        ('/assets/js/supabase-config.js', 'Capa 1: Core', 'Conexión única y centralizada con base de datos y autenticación mediante cliente oficial.'),
        ('/assets/js/lead-service.js', 'Capa 3: Lógica & Logs', 'Módulo universal de recepción, validación, telemetría y envío de leads con log centralizado (Sentry/LogSnag).'),
        ('/assets/js/vox-license.js', 'Capa 3: Lógica', 'Control de licenciamiento corporativo, validación de créditos y mecanismos de contingencia/killswitch.'),
        ('/servicios/*.html', 'Capa 2: Páginas', 'Páginas individuales de cada servicio (desarrollo, contable, jurídico, etc.) consumiendo el Core compartido.'),
        ('/admin/ (dashboard.html, .js)', 'Capa 4: Gestión', 'Panel de administración protegido por RLS, tokens JWT y control de acceso basado en roles (RBAC).'),
        ('/api/ (leads.js, verify-license.js)', 'Capa 4: Backend', 'Funciones serverless seguras protegidas con CORS, Rate Limiting por IP y validación de tokens anti-bot.')
    ]

    for r_idx, row_data in enumerate(files_data):
        row_cells = files_table.add_row().cells
        for i, val in enumerate(row_data):
            set_cell_background(row_cells[i], 'FFFFFF' if r_idx % 2 == 0 else 'F9FAFB')
            set_cell_margins(row_cells[i], top=60, bottom=60, left=80, right=80)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if i == 0:
                r.font.bold = True

    # -------------------------------------------------------------
    # 4. FLUJO DEV/OPS Y CONTROL DE ENTORNOS (NUEVA SECCIÓN ESTRATÉGICA)
    # -------------------------------------------------------------
    add_section_header("Flujo Dev/Ops, Gestión de Entornos y CI/CD", "4.")
    
    devops_items = [
        ("Estrategia de Control de Versiones (Git Workflow Estricto):", 
         "El repositorio sigue una arquitectura de branching protegida:\n"
         "• main: Rama exclusiva de Producción. Despliegue automático a producción tras pasar el pipeline de validación. Protegida contra push directo.\n"
         "• staging: Rama de Pruebas y Homologación (QA). Utilizada para validación de integración antes de release final.\n"
         "• feature/* o fix/*: Ramas efímeras de trabajo individual (ej. feature/turnstile-captcha, fix/cors-lead-api). Requieren Pull Request y revisión para integrarse."),
        
        ("Gestión Segura de Variables de Entorno (Envs):", 
         "Para garantizar cero fugas de credenciales en repositorios públicos/privados:\n"
         "• .env.example: Archivo versionado en Git con las llaves requeridas sin valores sensibles (ej. SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY=xxx).\n"
         "• .env.local: Archivo local de desarrollo excluido estrictamente en .gitignore.\n"
         "• .env.production / .env.staging: Variables inyectadas en tiempo de ejecución en la plataforma de hosting (Vercel Project Settings) con segregación estricta entre Service Role Keys (solo backend /api/) y Anon Public Keys (frontend)."),
        
        ("Pipeline de CI/CD e Integración Continua:", 
         "Automatización de pruebas y despliegues con GitHub Actions / Vercel Deploy Hooks:\n"
         "1. Trigger de PR: Ejecución automática de validación de sintaxis (Node syntax check), linter de código y verificación de dependencias.\n"
         "2. Preview Deployment: Creación automática de entorno de pruebas efímero para validar UI/UX y funcionalidad.\n"
         "3. Production Gate: El paso a 'main' requiere validación satisfactoria de las pruebas automatizadas y aprobación del Checklist de Homologación.")
    ]

    for title, desc in devops_items:
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(3)
        p_d.paragraph_format.space_after = Pt(4)
        r_t = p_d.add_run(f"4.{devops_items.index((title, desc)) + 1} {title}\n")
        r_t.font.name = 'Segoe UI'
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_PRIMARY
        
        r_b = p_d.add_run(desc)
        r_b.font.name = 'Segoe UI'
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # 5. PROTOCOLO DE CALIDAD DE CÓDIGO (QA) Y CONTROL DE ASSETS
    # -------------------------------------------------------------
    add_section_header("Protocolo de Pruebas, Calidad de Código (QA) y Manejo de Assets", "5.")
    
    qa_items = [
        ("Linter y Formateador Estándar:", 
         "Para mantener consistencia absoluta entre todos los desarrolladores del equipo:\n"
         "• ESLint & Prettier: Reglas de formateo unificadas (tabulaciones de 2 espacios, comillas simples, punto y coma obligatorio, sin variables huérfanas).\n"
         "• Stylelint: Validación de tokens CSS evitando el uso de valores hexadecimales 'hardcodeados' fuera de :root.\n"
         "• Pre-commit Hooks (Husky / Lint-Staged): Verificación obligatoria antes de permitir cualquier 'git commit'."),
        
        ("Manejo de Caché y Versionado de Assets (Cache-Busting):", 
         "• Control de versión en recursos estáticos mediante query params de versión (ej. styles.css?v=2.4) o nombrado con hash (ej. bundle.[hash].js).\n"
         "• Configuración de encabezados HTTP 'Cache-Control': Recursos inmutables (fuentes, logos e imágenes fijas) con cache prolongado (max-age=31536000, immutable); scripts y HTML con validación 'no-cache' o 'must-revalidate' para reflejar actualizaciones inmediatas sin requerir borrado manual de caché por el cliente.")
    ]

    for title, desc in qa_items:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(4)
        r_t = p_q.add_run(f"5.{qa_items.index((title, desc)) + 1} {title}\n")
        r_t.font.name = 'Segoe UI'
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_PRIMARY
        
        r_b = p_q.add_run(desc)
        r_b.font.name = 'Segoe UI'
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # 6. SEGURIDAD AMPLIADA (CAPA 4) & OBSERVABILIDAD (CAPA 3)
    # -------------------------------------------------------------
    add_section_header("Seguridad Ampliada (Capa 4), Privacidad y Observabilidad (Capa 3)", "6.")
    
    sec_obs_items = [
        ("Políticas CORS Estrictas y Rate Limiting Serverless:",
         "Todos los endpoints en /api/ (ej. leads.js) deben implementar cabeceras CORS que restrinjan el origen exclusivamente a dominios oficiales de VOX (voxbusinessdeveloper.com y previsualizaciones autorizadas). Asimismo, se implementa Rate Limiting por IP (máximo 5 peticiones de cotización por minuto por IP) para mitigar ataques DDoS, scrapping y spam masivo."),
        
        ("Protección Anti-Bots y Captcha Inteligente:",
         "Integración obligatoria en todos los formularios públicos (Capa 2/3) de Cloudflare Turnstile o reCAPTCHA v3 invisible. El token generado se valida en el backend serverless (/api/leads.js) previo al almacenamiento en Supabase o despacho de notificaciones vía Resend/Nodemailer."),
        
        ("Política de Cookies y Cumplimiento Normativo (GDPR / Privacidad):",
         "Implementación estándar de banner de consentimiento de cookies con categorización (Técnicas, Analíticas, Marketing), enlace permanente a la Política de Privacidad y Términos de Servicio, y bloqueo de scripts de tracking de terceros hasta que el usuario otorgue consentimiento explícito."),
        
        ("Monitoreo Centralizado de Log de Errores y Telemetría:",
         "Integración en lead-service.js y handlers globales de herramientas como Sentry o LogSnag. Cualquier fallo silencioso de red, error 500 de API o excepción no controlada en el cliente se reporta de forma asíncrona con stack trace y contexto (dispositivo, URL de servicio, timestamp) sin degradar la experiencia de usuario.")
    ]

    for title, desc in sec_obs_items:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(3)
        p_s.paragraph_format.space_after = Pt(4)
        r_t = p_s.add_run(f"6.{sec_obs_items.index((title, desc)) + 1} {title}\n")
        r_t.font.name = 'Segoe UI'
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_PRIMARY
        
        r_b = p_s.add_run(desc)
        r_b.font.name = 'Segoe UI'
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # 7. REGLAS TÉCNICAS OBLIGATORIAS PARA NUEVAS PÁGINAS
    # -------------------------------------------------------------
    add_section_header("Reglas Técnicas Obligatorias para Nuevas Páginas", "7.")

    rules = [
        ("1. Cero Código Duplicado (DRY Principle):", "Ninguna página debe reescribir lógica de envío de formularios, conexión con Supabase o configuración de APIs; siempre se importan modularmente desde /assets/js/."),
        ("2. Estilos Basados en Tokens Globales:", "Todos los colores, sombras, tipografías y radios de borde deben provenir de las variables maestras de CSS (:root) para garantizar coherencia y propagación inmediata."),
        ("3. Manejo Estandarizado de Formularios y Leads:", "Cada formulario debe emitir eventos unificados que capturen: servicio de origen, datos de contacto, requerimiento comercial, token anti-bot y metadatos de trazabilidad."),
        ("4. Arquitectura Accesible y Semántica (HTML5):", "Uso obligatorio de etiquetas semánticas (<header>, <main>, <section>, <article>, <aside>, <footer>) con un único <h1> por página, estructura jerárquica de encabezados (H2/H3) y atributos ARIA completos."),
        ("5. Optimización, Resiliencia y Feedback Visual:", "Manejo proactivo de errores (Try/Catch), feedback visual inmediato al usuario en cada acción (estados de carga con spinners, alertas toast de éxito/error) y prevención de render blocking.")
    ]

    for rule_title, rule_desc in rules:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_before = Pt(2)
        p_r.paragraph_format.space_after = Pt(3)
        r_rt = p_r.add_run(f"{rule_title} ")
        r_rt.font.name = 'Segoe UI'
        r_rt.font.size = Pt(9.5)
        r_rt.font.bold = True
        r_rt.font.color.rgb = COLOR_DARK
        
        r_rd = p_r.add_run(rule_desc)
        r_rd.font.name = 'Segoe UI'
        r_rd.font.size = Pt(9.5)
        r_rd.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # 8. CHECKLIST DE HOMOLOGACIÓN PRE-PUBLICACIÓN (ACTUALIZADO)
    # -------------------------------------------------------------
    add_section_header("Checklist de Homologación Pre-Publicación (Gate de Producción)", "8.")
    
    p_ch_intro = doc.add_paragraph()
    p_ch_intro.paragraph_format.space_after = Pt(4)
    r_ch = p_ch_intro.add_run("Antes de que cualquier página, servicio o módulo desarrollado por VOX sea publicado en el entorno de producción, debe superar y documentar el siguiente checklist técnico:")
    r_ch.font.name = 'Segoe UI'
    r_ch.font.size = Pt(9.5)

    checklist_table = doc.add_table(rows=1, cols=3)
    checklist_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    chk_hdr = checklist_table.rows[0].cells
    chk_hdr_titles = ['Estado', 'Área de Verificación', 'Criterio de Aceptación Obligatorio']
    for i, title in enumerate(chk_hdr_titles):
        set_cell_background(chk_hdr[i], '161A1F')
        set_cell_margins(chk_hdr[i], top=70, bottom=70, left=90, right=90)
        p = chk_hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_WHITE

    checks = [
        ('[ OK ]', 'Core & Arquitectura', 'Importación correcta del Core de estilos compartidos, tokens globales y scripts sin dependencias huérfanas.'),
        ('[ OK ]', 'Lead Service & Captcha', 'Formulario de contacto conectado con lead-service.js, protección anti-bot (Turnstile/reCAPTCHA) y registro en Supabase/Email.'),
        ('[ OK ]', 'DevOps & Variables (Envs)', 'Variables de entorno segregadas (.env.example actualizado, claves privadas configuradas en Vercel y ausentes en el frontend).'),
        ('[ OK ]', 'Calidad & Linters (QA)', 'Código validado sin errores de sintaxis AST, formateado bajo estándar ESLint/Prettier y con estrategia de cache-busting en assets.'),
        ('[ OK ]', 'Seguridad Backend & CORS', 'Endpoints /api/ con CORS restringido a dominios oficiales y Rate Limiting por IP configurado.'),
        ('[ OK ]', 'Monitoreo de Errores', 'Integración de logging de excepciones (Sentry/LogSnag) y manejo resiliente de fallos de red.'),
        ('[ OK ]', 'SEO & Accesibilidad Semántica', 'Metaetiquetas OpenGraph completas, Title descriptivo, Meta Description optimizada, único <h1> por página y jerarquía H2/H3.'),
        ('[ OK ]', 'Rendimiento (Performance)', 'Tiempos de carga menores a 1.5 - 2.0 segundos en conexiones móviles/estándar (Core Web Vitals verdes).'),
        ('[ OK ]', 'Responsive Design', 'Adaptabilidad móvil verificada en breakpoints estándar (360px, 768px, 1024px, 1440px) sin desbordamientos horizontales.'),
        ('[ OK ]', 'Consola Limpia & Privacidad', 'Consola del navegador sin advertencias o errores críticos y banner de consentimiento de cookies funcional.')
    ]

    for r_idx, check_row in enumerate(checks):
        row_cells = checklist_table.add_row().cells
        for i, val in enumerate(check_row):
            set_cell_background(row_cells[i], 'FFFFFF' if r_idx % 2 == 0 else 'F9FAFB')
            set_cell_margins(row_cells[i], top=50, bottom=50, left=70, right=70)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_SUCCESS
            elif i == 1:
                r.font.bold = True

    # -------------------------------------------------------------
    # 9. OPINIÓN Y DICTAMEN ARQUITECTÓNICO DEL SISTEMA
    # -------------------------------------------------------------
    add_section_header("Dictamen y Evaluación de la Arquitectura VOX", "9.")
    
    dictamen_box = (
        "1. Arquitectura Altamente Escalable:\n"
        "La división modular en 5 capas claras (Core, Componentes, Lógica de Servicios, Seguridad/Datos y UX/Rendimiento) "
        "garantiza un mantenimiento eficiente a largo plazo, desacopla el ciclo de vida del frontend y backend, "
        "y evita que la base de código se degrade con el tiempo ante el crecimiento del catálogo de servicios.\n\n"
        "2. Enfoque Orientado a Producto y Conversión Comercial:\n"
        "El hecho de estandarizar la captura de leads mediante un servicio universal (lead-service.js) desacoplado de las páginas individuales "
        "asegura que no se pierda información comercial en ningún punto del ecosistema, unificando trazabilidad, analítica y seguridad en un solo flujo."
    )
    
    add_callout(
        dictamen_box,
        prefix="EVALUACIÓN DE INGENIERÍA & IMPACTO EN EL NEGOCIO:",
        border_color="27AE60",
        bg_color="F0FFF4"
    )

    # Guardar en la raíz y en docs
    output_root = r"c:\Users\USUARIO\Desktop\Pagina web VOX\VOX_Manual_Arquitectura_Desarrollo_Web.docx"
    output_docs = r"c:\Users\USUARIO\Desktop\Pagina web VOX\docs\VOX_Manual_Arquitectura_Desarrollo_Web.docx"
    
    doc.save(output_root)
    doc.save(output_docs)
    print("Documento guardado exitosamente en:")
    print(f" - {output_root}")
    print(f" - {output_docs}")

if __name__ == "__main__":
    build_vox_architecture_manual()
